"""Evidence adapter for SMILES2Docking ligand 3D preparation and protonation."""

from __future__ import annotations

import contextlib
import json
import re
from collections.abc import Mapping
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from fidelichem.adapters.base import (
    build_import_plan,
    compute_source_artifacts,
    scan_source_files,
    verify_import_plan,
)
from fidelichem.adapters.gold.parsers import parse_gold_mol2
from fidelichem.chemistry.service import ChemistryService
from fidelichem.domain.adapters import (
    DetectionReport,
    ImportBundle,
    ImportPlan,
    QCIssue,
    QCSeverity,
    RawCompoundRecord,
    SourceArtifactRecord,
    ValidationReport,
)


def _base_structure_id(value: str) -> str:
    """Return the source molecule id from a generated stereoisomer id."""
    return value.split("__", 1)[0].strip()


def _normalise_header(value: object) -> str:
    """Normalise spreadsheet headers for tolerant identity matching."""
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def _load_external_input_smiles(
    source_root: Path,
    json_files: list[str],
) -> dict[str, str]:
    """Load source-id/SMILES pairs referenced by a SMILES2Docking report.

    The preparation workflow stores the original XLSX path in its JSON report,
    while the generated MOL2 contains only coordinates and source ids.  Reading
    that input table restores the chemical identity without guessing from a
    ligand filename.
    """
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError:
        return {}

    pairs: dict[str, str] = {}
    for json_rel in json_files:
        report_path = source_root / json_rel
        try:
            data = json.loads(report_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError, TypeError):
            continue
        if not isinstance(data, dict):
            continue
        input_value = data.get("input_file")
        if not isinstance(input_value, str) or not input_value.strip():
            continue
        input_path = Path(input_value)
        if not input_path.is_absolute():
            input_path = source_root / input_path
        if input_path.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        if not input_path.exists():
            continue
        try:
            workbook = openpyxl.load_workbook(
                filename=input_path,
                read_only=True,
                data_only=True,
            )
        except (OSError, ValueError, TypeError):
            continue
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            headers = [_normalise_header(value) for value in header_row]
            id_index = next(
                (
                    index
                    for index, header in enumerate(headers)
                    if header
                    in {
                        "accesscode",
                        "moleculeid",
                        "compoundid",
                        "ligandid",
                        "id",
                    }
                ),
                None,
            )
            smiles_index = next(
                (
                    index
                    for index, header in enumerate(headers)
                    if header
                    in {"smiles", "canonicalsmiles", "isomericsmiles"}
                ),
                None,
            )
            if id_index is None or smiles_index is None:
                continue
            for row in rows:
                if id_index >= len(row) or smiles_index >= len(row):
                    continue
                source_id = row[id_index]
                smiles = row[smiles_index]
                if source_id is None or smiles is None:
                    continue
                clean_id = str(source_id).strip()
                clean_smiles = str(smiles).strip()
                if clean_id and clean_smiles:
                    pairs.setdefault(clean_id, clean_smiles)
        except (OSError, ValueError, TypeError):
            continue
        finally:
            workbook.close()
    return pairs


def _derive_mol2_smiles(raw_block: str) -> str | None:
    """Derive a validated SMILES from a MOL2 block when no source table exists."""
    return ChemistryService.derive_smiles_from_mol2_block(raw_block)


class Smiles2DockingAdapter:
    """Evidence adapter parsing SMILES2Docking runs and protonation states."""

    adapter_id = "fidelichem.smiles2docking"
    adapter_version = "0.1.0"
    display_name = "SMILES2Docking Ligand Preparation Adapter"

    def probe(self, source: Path) -> DetectionReport:
        """Probe source directory for SMILES2Docking run descriptors."""
        if not source.exists():
            return DetectionReport(
                confidence=0.0,
                detected_format="unknown",
                candidate_files=(),
                requires_user_mapping=False,
                suggested_adapter=self.adapter_id,
            )

        patterns = (
            "*smiles2docking*.json",
            "run.json",
            "manifest.json",
            "run_report*.json",
            "prepared_*.sdf",
            "prepared_*.mol2",
            "*.sdf",
        )
        matches = scan_source_files(source, patterns=patterns)
        match_names = [f.as_posix() for f in matches]

        json_descriptors = [
            f
            for f in match_names
            if f.lower().endswith(".json")
            and (
                "smiles2docking" in f.lower()
                or "run_report" in f.lower()
                or f.lower() in ("run.json", "manifest.json")
            )
        ]

        if json_descriptors:
            return DetectionReport(
                confidence=0.95,
                detected_format="smiles2docking_run",
                candidate_files=tuple(match_names),
                requires_user_mapping=False,
                suggested_adapter=self.adapter_id,
            )

        sdf_or_mol2 = [
            f
            for f in match_names
            if "prepared" in f.lower() or f.endswith((".sdf", ".mol2"))
        ]
        if sdf_or_mol2:
            return DetectionReport(
                confidence=0.75,
                detected_format="smiles2docking_structures",
                candidate_files=tuple(match_names),
                requires_user_mapping=False,
                suggested_adapter=self.adapter_id,
            )

        return DetectionReport(
            confidence=0.0,
            detected_format="unknown",
            candidate_files=(),
            requires_user_mapping=False,
            suggested_adapter=self.adapter_id,
        )

    def plan(
        self,
        source: Path,
        options: Mapping[str, Any] | None = None,
    ) -> ImportPlan:
        """Build content-addressed import plan for SMILES2Docking evidence."""
        detection = self.probe(source)
        candidate_files = (
            list(detection.candidate_files)
            if detection.candidate_files
            else [p.as_posix() for p in scan_source_files(source)]
        )
        artifacts = compute_source_artifacts(source, candidate_files)
        return build_import_plan(
            adapter_id=self.adapter_id,
            adapter_version=self.adapter_version,
            source_root=str(source),
            artifacts=artifacts,
            options=options,
        )

    def parse(self, plan: ImportPlan) -> ImportBundle:
        """Parse prepared compounds, pH, tautomer states, and structure paths."""
        verify_import_plan(plan)
        source_root = Path(plan.source_root)
        source_artifacts: list[SourceArtifactRecord] = []
        qc_messages: list[QCIssue] = []
        compounds: list[RawCompoundRecord] = []

        file_hash_map = dict(plan.file_hashes)
        for rel_path in plan.source_files:
            file_path = source_root / rel_path
            sha = file_hash_map.get(rel_path, "")
            stat = file_path.stat() if file_path.exists() else None
            mtime = datetime.fromtimestamp(stat.st_mtime, tz=UTC) if stat else None
            size = stat.st_size if stat else 0
            file_type = (
                "json"
                if rel_path.lower().endswith(".json")
                else (
                    "sdf"
                    if rel_path.lower().endswith(".sdf")
                    else "mol2"
                )
            )

            source_artifacts.append(
                SourceArtifactRecord(
                    relative_path=rel_path,
                    sha256=sha,
                    size_bytes=size,
                    file_type=file_type,
                    mtime=mtime,
                )
            )

        json_files = [
            f
            for f in plan.source_files
            if f.lower().endswith(".json")
            and (
                "smiles2docking" in f.lower()
                or "run_report" in f.lower()
                or f.lower() in ("run.json", "manifest.json")
            )
        ]

        seen_cmpd_ids: set[str] = set()

        for json_rel in json_files:
            json_path = source_root / json_rel
            with contextlib.suppress(Exception):
                content = json_path.read_text(encoding="utf-8")
                data = json.loads(content)
                global_ph = (
                    data.get("target_ph")
                    or data.get("preparation_ph")
                    or data.get("ph")
                )
                global_ph_val = float(global_ph) if global_ph is not None else None
                prot_engine = data.get("protonation_engine") or data.get(
                    "protonation_backend"
                )
                opt_method = data.get("optimization_method") or data.get(
                    "energy_minimization"
                )

                items = (
                    data.get("compounds")
                    or data.get("ligands")
                    or data.get("results")
                    or []
                )
                if isinstance(data, list):
                    items = data

                for idx, item in enumerate(items, start=1):
                    if not isinstance(item, dict):
                        continue
                    cmpd_id = str(
                        item.get("compound_id")
                        or item.get("id")
                        or item.get("name")
                        or f"LIG_{idx}"
                    )
                    state_smiles = (
                        item.get("state_smiles")
                        or item.get("protonated_smiles")
                        or item.get("smiles")
                    )
                    initial_smiles = item.get("initial_smiles")
                    struct_file = (
                        item.get("structure_file")
                        or item.get("structure_path")
                        or item.get("file")
                    )

                    item_ph = item.get("ph")
                    eff_ph = float(item_ph) if item_ph is not None else global_ph_val

                    meta = {
                        "source_adapter": self.adapter_id,
                        "initial_smiles": initial_smiles,
                        "protonation_engine": (
                            prot_engine or item.get("protonation_engine")
                        ),
                        "optimization_method": (
                            opt_method or item.get("optimization_method")
                        ),
                        "structure_file": struct_file,
                        **{
                            k: v
                            for k, v in item.items()
                            if k
                            not in (
                                "compound_id",
                                "id",
                                "smiles",
                                "state_smiles",
                                "protonated_smiles",
                            )
                        },
                    }

                    eff_smiles = (
                        str(state_smiles)
                        if state_smiles
                        else (str(initial_smiles) if initial_smiles else None)
                    )

                    seen_cmpd_ids.add(cmpd_id)
                    compounds.append(
                        RawCompoundRecord(
                            source_system="smiles2docking",
                            source_value=cmpd_id,
                            source_smiles=eff_smiles,
                            preparation_ph=eff_ph,
                            source_artifact_path=struct_file or json_rel,
                            metadata=meta,
                        )
                    )

        identity_smiles = _load_external_input_smiles(source_root, json_files)
        mol2_files = [
            rel_path
            for rel_path in plan.source_files
            if rel_path.lower().endswith(".mol2")
        ]
        for mol2_rel in mol2_files:
            mol2_path = source_root / mol2_rel
            try:
                structures = parse_gold_mol2(mol2_path)
            except (OSError, UnicodeError, ValueError, TypeError) as exc:
                qc_messages.append(
                    QCIssue(
                        code="QC_MOL2_READ_FAILED",
                        message=f"Could not read MOL2 source: {exc}",
                        severity=QCSeverity.ERROR,
                        source_file=mol2_rel,
                    )
                )
                continue

            for structure in structures:
                compound_id = str(structure.get("ligand_name", "")).strip()
                if not compound_id or compound_id in seen_cmpd_ids:
                    continue
                base_id = _base_structure_id(compound_id)
                source_smiles = structure.get("smiles") or identity_smiles.get(
                    base_id
                )
                if not source_smiles:
                    source_smiles = _derive_mol2_smiles(
                        str(structure.get("raw_block", ""))
                    )
                if not source_smiles:
                    qc_messages.append(
                        QCIssue(
                            code="QC_MOL2_IDENTITY_MISSING",
                            message=(
                                "MOL2 structure has no embedded or referenced "
                                "SMILES identity"
                            ),
                            severity=QCSeverity.ERROR,
                            source_file=mol2_rel,
                            entity_reference=compound_id,
                        )
                    )

                seen_cmpd_ids.add(compound_id)
                compounds.append(
                    RawCompoundRecord(
                        source_system="smiles2docking",
                        source_value=compound_id,
                        source_smiles=(
                            str(source_smiles) if source_smiles else None
                        ),
                        source_artifact_path=mol2_rel,
                        metadata={
                            "source_adapter": self.adapter_id,
                            "structure_format": "mol2",
                            "atom_count": structure.get("atom_count", 0),
                            "bond_count": structure.get("bond_count", 0),
                            "base_structure_id": base_id,
                        },
                    )
                )

        return ImportBundle(
            plan=plan,
            targets=(),
            docking_runs=(),
            compounds=tuple(compounds),
            poses=(),
            scores=(),
            source_artifacts=tuple(source_artifacts),
            qc_messages=tuple(qc_messages),
            provenance={"adapter": self.adapter_id, "version": self.adapter_version},
        )

    def validate(self, bundle: ImportBundle) -> ValidationReport:
        """Validate parsed SMILES2Docking bundle."""
        errors: list[str] = []
        warnings: list[str] = []

        if not bundle.source_artifacts:
            errors.append("Bundle contains no source artifacts")

        if not bundle.compounds:
            errors.append(
                "No prepared ligands were parsed from SMILES2Docking evidence"
            )

        for qc in bundle.qc_messages:
            if qc.severity == QCSeverity.ERROR:
                errors.append(f"[{qc.code}] {qc.message}")
            elif qc.severity == QCSeverity.WARNING:
                warnings.append(f"[{qc.code}] {qc.message}")

        return ValidationReport(
            is_valid=len(errors) == 0,
            errors=tuple(errors),
            warnings=tuple(warnings),
            qc_issues=bundle.qc_messages,
        )


__all__ = ["Smiles2DockingAdapter"]
