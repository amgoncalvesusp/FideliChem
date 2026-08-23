"""Robust tabular file inspection, format detection, and streaming record extraction."""

from __future__ import annotations

import csv
import json
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from fidelichem.domain.table_importer import TableMappingSchema


def detect_format(file_path: Path) -> str:
    """Identify the tabular file format based on extension and content."""
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix == ".tsv":
        return "tsv"
    if suffix == ".json":
        return "json"
    if suffix in (".jsonl", ".ndjson"):
        return "jsonl"
    if suffix == ".parquet":
        return "parquet"
    if suffix in (".xlsx", ".xlsm"):
        return "xlsx"
    if suffix == ".xls":
        return "xls"

    # Fallback to content inspection if extension is generic
    try:
        sample = file_path.read_text(encoding="utf-8", errors="ignore")[:4096].strip()
        if sample.startswith("[") and sample.endswith("]"):
            return "json"
        if sample.startswith("{"):
            return "jsonl"
    except Exception:
        pass

    return "unknown"


def detect_delimiter(file_path: Path) -> str:
    """Inspect CSV/TSV header and content to deduce the field delimiter."""
    suffix = file_path.suffix.lower()
    if suffix == ".tsv":
        return "\t"

    try:
        sample = file_path.read_text(encoding="utf-8", errors="ignore")[:8192]
        lines = [
            line.strip()
            for line in sample.splitlines()
            if line.strip() and not line.startswith("#")
        ]
        if not lines:
            return ","

        first_line = lines[0]
        # Count candidate delimiter frequencies in first line
        delimiters = [",", "\t", ";", "|"]
        counts = {d: first_line.count(d) for d in delimiters}
        best_delimiter = max(counts, key=counts.get)  # type: ignore[arg-type]
        if counts[best_delimiter] > 0:
            return best_delimiter

        # Try csv.Sniffer fallback
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff("\n".join(lines[:5]), delimiters=",\t;|")
        return dialect.delimiter
    except Exception:
        return ","


def preview_table(
    file_path: Path,
    max_rows: int = 10,
    *,
    delimiter: str | None = None,
    sheet_name: str | None = None,
) -> tuple[tuple[str, ...], tuple[tuple[Any, ...], ...]]:
    """Return headers and preview rows for UI and mapping setup."""
    fmt = detect_format(file_path)

    if fmt in ("csv", "tsv"):
        delim = delimiter or detect_delimiter(file_path)
        with file_path.open(
            mode="r", encoding="utf-8", errors="replace", newline=""
        ) as f:
            reader = csv.reader(f, delimiter=delim)
            rows: list[list[str]] = []
            for row in reader:
                if not row or (row and row[0].startswith("#")):
                    continue
                rows.append([cell.strip() for cell in row])
                if len(rows) > max_rows:
                    break

            if not rows:
                return ((), ())

            headers = tuple(rows[0])
            csv_data_rows: tuple[tuple[Any, ...], ...] = tuple(
                tuple(r) for r in rows[1 : max_rows + 1]
            )
            return (headers, csv_data_rows)

    if fmt == "json":
        with file_path.open(mode="r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)
            if isinstance(data, list) and data and isinstance(data[0], dict):
                headers = tuple(sorted(data[0].keys()))
                json_data_rows: list[tuple[Any, ...]] = []
                for item in data[:max_rows]:
                    if isinstance(item, dict):
                        json_data_rows.append(tuple(item.get(h) for h in headers))
                return (headers, tuple(json_data_rows))
            return ((), ())

    if fmt == "jsonl":
        with file_path.open(mode="r", encoding="utf-8", errors="replace") as f:
            records: list[dict[str, Any]] = []
            for line in f:
                stripped = line.strip()
                if stripped and not stripped.startswith("#"):
                    try:
                        obj = json.loads(stripped)
                        if isinstance(obj, dict):
                            records.append(obj)
                            if len(records) >= max_rows + 1:
                                break
                    except json.JSONDecodeError:
                        continue

            if not records:
                return ((), ())

            all_keys: set[str] = set()
            for r in records:
                all_keys.update(r.keys())
            headers = tuple(sorted(all_keys))

            jsonl_data_rows: tuple[tuple[Any, ...], ...] = tuple(
                tuple(r.get(h) for h in headers) for r in records[:max_rows]
            )
            return (headers, jsonl_data_rows)

    if fmt in ("xlsx", "xlsm", "xls", "parquet"):
        records = list(
            read_table_records(
                file_path,
                TableMappingSchema(sheet_name=sheet_name),
            )
        )[:max_rows]
        if not records:
            return ((), ())
        headers = tuple(records[0].keys())
        preview_rows = tuple(
            tuple(record.get(header) for header in headers) for record in records
        )
        return headers, preview_rows

    return ((), ())


def _clean_cell(val: Any) -> Any:
    """Normalize cell values: empty strings and whitespace become None."""
    if val is None:
        return None
    if isinstance(val, str):
        cleaned = val.strip()
        return cleaned if cleaned else None
    return val


def read_table_records(
    file_path: Path,
    schema: TableMappingSchema,
) -> Iterator[dict[str, Any]]:
    """Stream tabular rows as normalized dictionaries of column -> value."""
    fmt = detect_format(file_path)

    if fmt in ("csv", "tsv"):
        delim = schema.delimiter or detect_delimiter(file_path)
        with file_path.open(
            mode="r", encoding="utf-8", errors="replace", newline=""
        ) as f:
            # Skip initial rows if requested
            for _ in range(schema.skip_rows):
                f.readline()

            reader = csv.reader(f, delimiter=delim)
            headers: list[str] = []

            for row in reader:
                if not row:
                    continue
                if schema.comment_prefix and row[0].startswith(schema.comment_prefix):
                    continue

                if not headers:
                    if schema.has_header:
                        headers = [cell.strip() for cell in row]
                        continue
                    headers = [f"col_{i + 1}" for i in range(len(row))]

                record: dict[str, Any] = {}
                for idx, col_name in enumerate(headers):
                    raw_val = row[idx] if idx < len(row) else None
                    record[col_name] = _clean_cell(raw_val)

                yield record

    elif fmt == "json":
        with file_path.open(mode="r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)
            if isinstance(data, list):
                for item in data:
                    if isinstance(item, dict):
                        yield {k: _clean_cell(v) for k, v in item.items()}

    elif fmt == "jsonl":
        with file_path.open(mode="r", encoding="utf-8", errors="replace") as f:
            for line in f:
                stripped = line.strip()
                if stripped and (
                    not schema.comment_prefix
                    or not stripped.startswith(schema.comment_prefix)
                ):
                    try:
                        obj = json.loads(stripped)
                        if isinstance(obj, dict):
                            yield {k: _clean_cell(v) for k, v in obj.items()}
                    except json.JSONDecodeError:
                        continue

    elif fmt in ("xlsx", "xlsm"):
        try:
            import openpyxl  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ValueError(
                "XLSX input requires the optional 'openpyxl' dependency"
            ) from exc

        workbook = openpyxl.load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )
        try:
            sheet_name = schema.sheet_name or workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet '{sheet_name}' was not found")
            rows = workbook[sheet_name].iter_rows(values_only=True)
            for _ in range(schema.skip_rows):
                next(rows, None)
            sheet_headers: list[str] = []
            for raw_row in rows:
                row = list(raw_row)
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                if not sheet_headers:
                    if schema.has_header:
                        sheet_headers = [str(value).strip() for value in row]
                        continue
                    sheet_headers = [f"col_{i + 1}" for i in range(len(row))]
                yield {
                    header: _clean_cell(row[index] if index < len(row) else None)
                    for index, header in enumerate(sheet_headers)
                }
        finally:
            workbook.close()

    elif fmt == "xls":
        try:
            import xlrd  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ValueError(
                "XLS input requires the optional 'xlrd' dependency"
            ) from exc

        try:
            workbook = xlrd.open_workbook(filename=str(file_path), on_demand=True)
        except Exception as exc:  # noqa: BLE001
            raise ValueError("XLS input could not be read") from exc
        try:
            sheet_name = schema.sheet_name
            if sheet_name is None:
                sheet = workbook.sheet_by_index(0)
            elif sheet_name in workbook.sheet_names:
                sheet = workbook.sheet_by_name(sheet_name)
            else:
                raise ValueError(f"Worksheet '{sheet_name}' was not found")

            xls_headers: list[str] = []
            for row_index in range(schema.skip_rows, sheet.nrows):
                row = list(sheet.row_values(row_index))
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                if not xls_headers:
                    if schema.has_header:
                        xls_headers = [str(value).strip() for value in row]
                        continue
                    xls_headers = [f"col_{i + 1}" for i in range(len(row))]
                yield {
                    header: _clean_cell(row[index] if index < len(row) else None)
                    for index, header in enumerate(xls_headers)
                }
        finally:
            workbook.release_resources()

    elif fmt == "parquet":
        try:
            import pyarrow.parquet as parquet  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ValueError(
                "Parquet input requires the optional 'pyarrow' dependency"
            ) from exc
        table = parquet.read_table(file_path)
        for item in table.to_pylist():
            if isinstance(item, dict):
                yield {key: _clean_cell(value) for key, value in item.items()}


__all__ = [
    "detect_delimiter",
    "detect_format",
    "preview_table",
    "read_table_records",
]
